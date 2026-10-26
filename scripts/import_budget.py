"""Convertit le classeur Excel « Budget » en fichier JSON à importer dans Finance > Modèle.

Usage : python scripts/import_budget.py "chemin/vers/Budget 2026.xlsx" [sortie.json] [--annee 2026]
Par défaut, le JSON est écrit à côté du classeur. Il contient tes montants : ne le mets pas dans le dépôt.

L'onglet TYPE devient le modèle, chaque onglet de mois (JUIN, JUILLET…) un mois.
Une cellule comme « =7.9+10.9+6 » est décomposée en autant de montants, pour garder le détail.
"""
import json, re, sys
from pathlib import Path
import openpyxl

args = [a for a in sys.argv[1:] if not a.startswith("--")]
XLSX = Path(args[0])
OUT = Path(args[1]) if len(args) > 1 else XLSX.with_name(XLSX.stem + " - FSVS.json")
year_arg = sys.argv[sys.argv.index("--annee") + 1] if "--annee" in sys.argv else None
YEAR = int(year_arg or (re.search(r"20\d\d", XLSX.stem) or [None])[0] or 0)
if not YEAR:
    sys.exit("Année introuvable dans le nom du fichier : ajoute --annee 2026")

MONTHS = ["JANVIER", "FEVRIER", "MARS", "AVRIL", "MAI", "JUIN", "JUILLET", "AOUT", "SEPTEMBRE", "OCTOBRE", "NOVEMBRE", "DECEMBRE"]
# Postes : (id, libellé, icône, cellule du titre, cellule de la cible)
GROUPS = [
    ("transports", "Transports", "🚇", "F17", "G35"),
    ("nourriture", "Nourriture", "🛒", "I17", "J35"),
    ("abonnements", "Abonnements", "🔁", "F22", "G38"),
    ("divers", "Dépenses diverses", "🛍️", "I22", "J38"),
    ("investissements", "Investissements", "📈", "L17", "M35"),
]

wb_f = openpyxl.load_workbook(XLSX)
wb_v = openpyxl.load_workbook(XLSX, data_only=True)
counter = 0


def uid(prefix):
    global counter
    counter += 1
    return f"{prefix}-x{counter}"


def slug(s):
    s = s.lower()
    for a, b in (("é", "e"), ("è", "e"), ("ê", "e"), ("ô", "o"), ("â", "a"), ("ç", "c")):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def amounts(ws_f, ws_v, coord):
    """Montants d'une cellule : « =8.03+8.48 » donne [8.03, 8.48] ; une formule avec références donne sa valeur."""
    f, v = ws_f[coord].value, ws_v[coord].value
    if isinstance(f, str) and re.fullmatch(r"=[\d.\s+\-]+", f):
        terms = re.findall(r"([+-]?)\s*(\d+(?:\.\d+)?)", f[1:])
        return [round(float(n) * (-1 if s == "-" else 1), 2) for s, n in terms if float(n)]
    return [round(float(v), 2)] if isinstance(v, (int, float)) and v else []


def lines_below(ws_f, ws_v, title):
    """Lignes sous un titre de poste, jusqu'à « Total »."""
    col, row = openpyxl.utils.cell.coordinate_from_string(title)
    ci = openpyxl.utils.column_index_from_string(col)
    out, r = [], row + 1
    while (label := ws_v.cell(r, ci).value) and str(label).strip() != "Total":
        coord = ws_v.cell(r, ci + 1).coordinate
        out.append({"id": slug(str(label)), "label": str(label).strip(),
                    "ops": [{"id": uid("op"), "amount": a} for a in amounts(ws_f, ws_v, coord)]})
        r += 1
    return out


def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else None


def read_month(name):
    ws_f, ws_v = wb_f[name], wb_v[name]
    incomes = []
    for r in range(4, 9):
        label = ws_v.cell(r, 2).value
        if label:
            incomes.append({"id": slug(str(label)), "label": str(label).strip(),
                            "ops": [{"id": uid("op"), "amount": a} for a in amounts(ws_f, ws_v, f"D{r}")]})
    groups = [{"id": gid, "label": label, "icon": icon, "target": num(ws_v[target].value) or 0,
               "lines": lines_below(ws_f, ws_v, title)} for gid, label, icon, title, target in GROUPS]
    month = {"incomes": incomes, "groups": groups}
    if (opening := num(ws_v["B18"].value)) is not None:
        month["opening"] = opening
    # Solde réel : « MONTANT REEL » (O21), sinon un montant saisi à la main en O18 (pas la formule de solde prévu).
    closing = num(ws_v["O21"].value)
    if closing is None and not str(ws_f["O18"].value or "").startswith("="):
        closing = num(ws_v["O18"].value)
    if closing is not None:
        month["closing"] = closing
    return month


template = read_month("TYPE")
template.pop("opening", None)
template.pop("closing", None)
months = {}
for name in wb_f.sheetnames:
    if name.upper() in MONTHS:
        months[f"{YEAR}-{MONTHS.index(name.upper()) + 1:02d}"] = read_month(name)

# Vérification : les totaux doivent retomber sur ceux du classeur.
for name in wb_v.sheetnames:
    if name.upper() not in MONTHS:
        continue
    m, ws = months[f"{YEAR}-{MONTHS.index(name.upper()) + 1:02d}"], wb_v[name]
    total = lambda ls: round(sum(o["amount"] for l in ls for o in l["ops"]), 2)
    income, spent = total(m["incomes"]), round(sum(total(g["lines"]) for g in m["groups"]), 2)
    excel_income, excel_spent = num(ws["D9"].value), num(ws["O10"].value)
    ok = income == excel_income and spent == excel_spent
    print(f"{name:<10} revenus {income:>9.2f}  dépenses {spent:>9.2f}  économies {income - spent:>9.2f}"
          + ("" if ok else f"   ≠ Excel {excel_income} / {excel_spent}"))

OUT.write_text(json.dumps({"template": template, "months": months}, ensure_ascii=False, indent=1))
print(f"→ {OUT}")
