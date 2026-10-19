"""Convertit le classeur Excel « Nutrition » en src/seed.json pour la webapp.

Usage : python scripts/import_excel.py "chemin/vers/Nutrition 2.xlsx" [--avec-suivi]
Par défaut le suivi part de zéro ; --avec-suivi reprend aussi l'historique.
Chaque journée du suivi (formule Excel) est décomposée en entrées
ingrédient / recette, puis recalculée et comparée aux valeurs d'Excel.
"""
import json, re, sys, datetime
from pathlib import Path
import openpyxl

XLSX = sys.argv[1]
OUT = Path(__file__).resolve().parent.parent / "src" / "seed.json"
NUTRIENTS = ["kcal", "proteines", "glucides", "sucres", "lipides", "satures", "fibres", "sel"]
ING_SHEET, REC_SHEET, DAY_SHEET = "Valeurs nutritionnelles par pro", "Repas", "Suivi KcalProtéines"
# Ingrédients dont les valeurs sont données pour une portion et non pour 100 g
PORTION_ROWS = {36: "scoop", 78: "barre", 87: "portion", 91: "portion"}

wb_f = openpyxl.load_workbook(XLSX)
wb_v = openpyxl.load_workbook(XLSX, data_only=True)


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


# ---- Ingrédients
ws = wb_v[ING_SHEET]
ingredients, ing_rows = [], {}
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 1).value
    if not name:
        continue
    j = ws.cell(r, 10).value
    ing = {
        "id": f"ing-{r}",
        "name": str(name).strip(),
        "basis": "portion" if r in PORTION_ROWS else "100g",
        "values": [round(num(ws.cell(r, c).value), 4) for c in range(2, 10)],
    }
    if r in PORTION_ROWS:
        ing["unitLabel"] = PORTION_ROWS[r]
    elif isinstance(j, (int, float)):
        ing["unitWeight"] = round(j * 100, 2)
    ingredients.append(ing)
    ing_rows[r] = ing

# ---- Recettes (valeurs du tableau + liste d'ingrédients quand elle existe)
ws = wb_v[REC_SHEET]
recipes, rec_rows = [], {}
for r in range(3, 21):
    name = ws.cell(r, 2).value
    if not name:
        continue
    rec = {
        "id": f"rec-{r}",
        "name": str(name).strip(),
        "portions": 1,
        "items": [],
        "manualValues": [round(num(ws.cell(r, c).value), 4) for c in range(3, 11)],
        "notes": "",
    }
    recipes.append(rec)
    rec_rows[r] = rec

NAME_MAP = {
    "riz": 2, "bavette": 3, "poivron": 6, "onion": 7, "oignon": 7, "échalotte": 8, "farine": 28,
    "sucre": 29, "pâtes": 5, "filet de poulet": 20, "poulet": 20, "aiguillettes poulet": 20,
    "crème semi épaisse": 18, "concentré de tomate": 19, "concombre": 9, "concombrer": 9,
    "skyr": 4, "skryr": 4, "mayonnaise": 11, "thon nature": 10, "saumon": 31, "beurre": 35,
    "fromage râpé": 32, "whey": 36, "sirop d'érable": 49, "beurre de cacahuète": 52,
    "graines de chia": 71, "flocons d'avoine": 44, "lait semi-écrémé": 45, "banane": 15,
    "avocat": 60, "oeuf": 14, "corn flakes": 50, "maïs": 61, "pomme de terre": 67,
    "haricots": 73, "haricots verts": 73,
}
by_name = {rec["name"].lower(): rec for rec in recipes}
for top in (22, 33, 41):
    for col in (2, 5, 8, 11):
        title = ws.cell(top, col).value
        if not title or title.lower() not in by_name:
            continue
        rec, unmapped = by_name[title.lower()], []
        r = top + 1
        while True:
            label, qty = ws.cell(r, col).value, ws.cell(r, col + 1).value
            if label is None:
                break
            label = str(label).strip()
            m = re.match(r"(\d+)\s*rep", label)
            if m:
                rec["portions"] = int(m.group(1))
                break
            r += 1
            if qty in (None, "NA"):
                continue
            row = NAME_MAP.get(label.lower())
            if row is None:
                unmapped.append(f"{label} {qty}")
                continue
            ing = ing_rows[row]
            if isinstance(qty, (int, float)):
                if ing.get("unitWeight"):
                    item = {"refId": ing["id"], "qty": qty, "unit": "unit"}
                else:
                    item = {"refId": ing["id"], "qty": qty * 100, "unit": "g"}
            else:
                q = str(qty).lower()
                n = float(re.match(r"[\d.]+", q).group(0))
                if "kg" in q:
                    item = {"refId": ing["id"], "qty": n * 1000, "unit": "g"}
                elif "scoop" in q or ing["basis"] == "portion":
                    item = {"refId": ing["id"], "qty": n, "unit": "portion"}
                else:
                    item = {"refId": ing["id"], "qty": n, "unit": "g"}
            rec["items"].append(item)
        if unmapped:
            rec["notes"] = "Non trouvés dans la base : " + ", ".join(unmapped)

# ---- Suivi quotidien
REF_ING = re.compile(r"ING!\$?([A-Z]+)\$?(\d+)")
REF_REC = re.compile(r"Repas!\$?([A-Z]+)\$?(\d+)")
wsf, wsv = wb_f[DAY_SHEET], wb_v[DAY_SHEET]
days, report = {}, []


def entry_values(e):
    if e["kind"] == "custom":
        return e["values"]
    if e["kind"] == "recipe":
        return [v * e["qty"] for v in rec_rows[int(e["refId"][4:])]["manualValues"]]
    ing = ing_rows[int(e["refId"][4:])]
    if e["unit"] == "g":
        f = e["qty"] / 100
    elif e["unit"] == "unit":
        f = e["qty"] * ing["unitWeight"] / 100
    else:
        f = e["qty"]
    return [v * f for v in ing["values"]]


def ingredient_entry(ing, coef, unit_ref, literals):
    uw = ing.get("unitWeight")
    if ing["basis"] == "portion":
        return {"qty": round(coef, 4), "unit": "portion"}
    if uw and unit_ref:
        return {"qty": round(coef, 4), "unit": "unit"}
    # ex. 3*B14*0.6 : nombre d'unités × poids unitaire écrit en dur
    if uw and len(literals) == 2 and any(abs(n - uw / 100) < 1e-9 for n in literals):
        return {"qty": round(coef / (uw / 100), 4), "unit": "unit"}
    return {"qty": round(coef * 100, 2), "unit": "g"}


for r in range(2, wsf.max_row + 1):
    d = wsf.cell(r, 1).value
    if not isinstance(d, datetime.datetime):
        continue
    f = wsf.cell(r, 2).value
    if f is None or f == "NA":
        continue
    cached = [wsv.cell(r, c).value for c in range(2, 10)]
    key = d.date().isoformat()
    entries = []
    if isinstance(f, (int, float)):
        entries.append({"kind": "custom", "label": "Saisie manuelle (Excel)", "values": [round(num(v), 4) for v in cached]})
    else:
        expr = f.lstrip("=").replace(f"'{ING_SHEET}'", "ING").replace(" ", "")
        for sign, term in re.findall(r"([+-]?)([^+-]+)", expr):
            coef, ref, unit_ref, literals = -1.0 if sign == "-" else 1.0, None, False, []
            for fac in term.split("*"):
                mi, mr = REF_ING.fullmatch(fac), REF_REC.fullmatch(fac)
                if mi and mi.group(1) == "B":
                    ref = ("ing", int(mi.group(2)))
                elif mi and mi.group(1) == "J":
                    unit_ref = True
                elif mr and mr.group(1) == "C":
                    ref = ("rec", int(mr.group(2)))
                elif mi or mr:
                    ref = ("skip", fac)  # cellule hors tableau (vide) : ignorée
                else:
                    coef *= float(fac)
                    literals.append(float(fac))
            if ref is None or ref[0] == "skip":
                continue
            if ref[0] == "rec":
                entries.append({"kind": "recipe", "refId": f"rec-{ref[1]}", "qty": round(coef, 4)})
            else:
                ing = ing_rows[ref[1]]
                entries.append({"kind": "ingredient", "refId": ing["id"], **ingredient_entry(ing, coef, unit_ref, literals)})
    for i, e in enumerate(entries):
        e["id"] = f"e-{key}-{i}"
    days[key] = entries
    total = [sum(col) for col in zip(*[entry_values(e) for e in entries])]
    diffs = [(NUTRIENTS[i], round(total[i], 2), round(num(cached[i]), 2)) for i in range(8) if abs(total[i] - num(cached[i])) > 0.05]
    report.append((key, len(entries), diffs))

objectives = {
    "kcal": {"target": 2400, "min": 2200, "max": 2650, "below": "warn", "above": "alert"},
    "proteines": {"target": 130, "min": 120, "max": 155, "below": "alert", "above": "warn"},
    "glucides": {"target": 300, "min": 270, "max": 350, "below": "warn", "above": "alert"},
    "sucres": {"max": 100, "below": "ok", "above": "alert"},
    "lipides": {"target": 70, "min": 60, "max": 80, "below": "warn", "above": "alert"},
    "satures": {},
    "fibres": {"target": 20, "min": 17, "below": "warn", "above": "ok"},
    "sel": {},
}

# Corrections de noms, partagées avec la migration de l'app (src/renames.json)
renames = json.loads((OUT.parent / "renames.json").read_text(encoding="utf-8"))
for items, key in ((ingredients, "ingredients"), (recipes, "recipes")):
    for x in items:
        if x["id"] in renames[key]:
            x["name"] = renames[key][x["id"]][1]

if "--avec-suivi" not in sys.argv:
    days = {}
OUT.write_text(json.dumps({"version": 3, "ingredients": ingredients, "recipes": recipes, "days": days, "objectives": objectives}, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"{len(ingredients)} ingrédients, {len(recipes)} recettes, {len(days)} jours -> {OUT}")
for key, n, diffs in report:
    print(key, n, "OK" if not diffs else diffs)
